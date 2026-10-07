"""Convert the NIST CSF 2.0 Reference Tool export (OLIR download) to compact JSON.

Source: https://csrc.nist.gov/extensions/nudp/services/json/csf/download?olirids=all
(sheet "CSF 2.0"; columns Function, Category, Subcategory, Implementation
Examples, Informative References).

Only two informative-reference families are kept, both as untyped "related"
mappings (the source states no subset/superset relationship):
  - "SP 800-53 Rev 5.2.0: <id>"  -> canonical control ids ("AC-2", "AC-2(3)")
  - "ISO/IEC 27001:2022: ..."    -> ids only (the ISO text is copyrighted), in ISO's
                                    notation: "Clause 4.1" (mandatory clauses 4-10) and
                                    "A.5.24" (Annex A controls). The two number spaces
                                    overlap (clause 6.1 vs Annex A 6.1), so the kind is kept.
The Rev 5.1.1 lines duplicate the 5.2.0 ones and are ignored, as is every other
framework. Withdrawn CSF 1.1 rows are skipped.

Usage:
  uv run python -m tools.nist_csf.convert_olir_xlsx \
      --xlsx olir_csf_all.xlsx --out data/nist/csf2_olir.json
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections.abc import Iterable, Sequence
from dataclasses import asdict, dataclass, field
from pathlib import Path

FRAMEWORK = "NIST CSF 2.0"
SOURCE_URL = "https://csrc.nist.gov/extensions/nudp/services/json/csf/download?olirids=all"
MAPPING_SOURCE = "NIST CSF 2.0 Reference Tool (informative references)"
SHEET = "CSF 2.0"

_NIST_PREFIX = "SP 800-53 Rev 5.2.0:"
_ISO_PREFIX = "ISO/IEC 27001:2022:"
_HEADER = ("Function", "Category", "Subcategory")

_FUNCTION_RE = re.compile(r"^(?P<name>.+?)\s*\((?P<id>[A-Z]{2})\):\s*(?P<text>.+)$", re.DOTALL)
_CATEGORY_RE = re.compile(r"^(?P<name>.+?)\s*\((?P<id>[A-Z]{2}\.[A-Z]{2})\):\s*(?P<text>.+)$", re.DOTALL)
_SUBCATEGORY_RE = re.compile(r"^(?P<id>[A-Z]{2}\.[A-Z]{2}-\d{2}):\s*(?P<text>.+)$", re.DOTALL)
_NIST_ID_RE = re.compile(r"^([A-Za-z]{2})-(\d+)(?:\((\d+)\))?$")
_ISO_KIND_RE = re.compile(r"^(?P<kind>Mandatory Clause|Annex A Controls|Control)\s*:?\s*(?P<body>.*)$")
_ISO_ID_RE = re.compile(r"^\d+(?:\.\d+)*(?:\([a-z]\))?$")


@dataclass
class CsfFunction:
    id: str
    name: str
    text: str


@dataclass
class CsfCategory:
    id: str
    function_id: str
    name: str
    text: str


@dataclass
class CsfSubcategory:
    id: str
    category_id: str
    text: str
    examples: list[str] = field(default_factory=list)
    refs_800_53: list[str] = field(default_factory=list)
    refs_iso_27001_2022: list[str] = field(default_factory=list)


@dataclass
class CsfCatalog:
    functions: list[CsfFunction] = field(default_factory=list)
    categories: list[CsfCategory] = field(default_factory=list)
    subcategories: list[CsfSubcategory] = field(default_factory=list)


def _cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _is_withdrawn(text: str) -> bool:
    return "[Withdrawn" in text


def canonical_nist_id(raw: str) -> str | None:
    """Canonical SP 800-53 id ("CM-07(02)" -> "CM-7(2)"); None for family-only refs like "PT"."""
    m = _NIST_ID_RE.match(raw.strip())
    if not m:
        return None
    enhancement = f"({int(m.group(3))})" if m.group(3) is not None else ""
    return f"{m.group(1).upper()}-{int(m.group(2))}{enhancement}"


def parse_iso_ids(line: str) -> list[str]:
    """Kind-qualified ids from one "ISO/IEC 27001:2022: ..." line (ids only, no ISO text).

    "Mandatory Clause: 4.2 (a)" -> ["Clause 4.2(a)"]; "Annex A Controls: 5.1, 5.2" and
    "Control  8.6" -> ["A.5.1", "A.5.2"], ["A.8.6"]. Empty and literal-"None" values and
    lines whose kind is unknown are skipped.
    """
    rest = line[len(_ISO_PREFIX) :].strip()
    kind = _ISO_KIND_RE.match(rest)
    if kind is None:
        return []
    prefix = "Clause " if kind.group("kind") == "Mandatory Clause" else "A."
    ids: list[str] = []
    for token in kind.group("body").split(","):
        compact = re.sub(r"\s+", "", token)
        if compact and compact != "None" and _ISO_ID_RE.match(compact):
            ids.append(f"{prefix}{compact}")
    return ids


def _nist_key(control_id: str) -> tuple[str, int, int]:
    m = _NIST_ID_RE.match(control_id)
    if m is None:
        raise ValueError(f"not a canonical control id: {control_id!r}")
    return (m.group(1), int(m.group(2)), int(m.group(3) or 0))


def _iso_key(iso_id: str) -> tuple[bool, tuple[int, ...], str]:
    return (iso_id.startswith("A."), tuple(int(n) for n in re.findall(r"\d+", iso_id)), iso_id)


def parse_csf_rows(rows: Iterable[Sequence[object]]) -> CsfCatalog:
    """Parse sheet rows (header and any title rows are tolerated) into a catalog."""
    catalog = CsfCatalog()
    seen_functions: set[str] = set()
    seen_categories: set[str] = set()

    for row in rows:
        cells = [*row, None, None, None, None, None][:5]
        function, category, subcategory, examples, refs = (_cell(c) for c in cells)
        if (function, category, subcategory) == _HEADER:
            continue

        if function:
            m = _FUNCTION_RE.match(function)
            if m and not _is_withdrawn(function) and m["id"] not in seen_functions:
                seen_functions.add(m["id"])
                catalog.functions.append(CsfFunction(m["id"], m["name"].strip(), m["text"].strip()))
        elif category:
            m = _CATEGORY_RE.match(category)
            if m and not _is_withdrawn(category) and m["id"] not in seen_categories:
                seen_categories.add(m["id"])
                catalog.categories.append(
                    CsfCategory(m["id"], m["id"].split(".")[0], m["name"].strip(), m["text"].strip())
                )
        elif subcategory:
            m = _SUBCATEGORY_RE.match(subcategory)
            if not m or _is_withdrawn(subcategory):
                continue
            sub_id = m["id"]
            nist: set[str] = set()
            iso: set[str] = set()
            for line in refs.splitlines():
                line = line.strip()
                if line.startswith(_NIST_PREFIX):
                    canonical = canonical_nist_id(line[len(_NIST_PREFIX) :])
                    if canonical:
                        nist.add(canonical)
                elif line.startswith(_ISO_PREFIX):
                    iso.update(parse_iso_ids(line))
            catalog.subcategories.append(
                CsfSubcategory(
                    id=sub_id,
                    category_id=sub_id.split("-")[0],
                    text=m["text"].strip(),
                    examples=[
                        re.sub(r"^Ex\d+:\s*", "", ex.strip()) for ex in examples.splitlines() if ex.strip()
                    ],
                    refs_800_53=sorted(nist, key=_nist_key),
                    refs_iso_27001_2022=sorted(iso, key=_iso_key),
                )
            )

    unknown = {s.category_id for s in catalog.subcategories} - seen_categories
    if unknown:
        raise ValueError(f"subcategories reference unknown categories: {sorted(unknown)}")
    return catalog


def catalog_to_json(catalog: CsfCatalog, *, source_sha256: str) -> dict[str, object]:
    return {
        "framework": FRAMEWORK,
        "source_url": SOURCE_URL,
        "source_sha256": source_sha256,
        "mapping_source": MAPPING_SOURCE,
        **asdict(catalog),
    }


def load_rows(xlsx: Path) -> list[tuple[object, ...]]:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx, read_only=True)
    try:
        return list(workbook[SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--out", default=Path("data/nist/csf2_olir.json"), type=Path)
    args = ap.parse_args()

    sha = hashlib.sha256(args.xlsx.read_bytes()).hexdigest()
    catalog = parse_csf_rows(load_rows(args.xlsx))
    payload = catalog_to_json(catalog, source_sha256=sha)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(
        json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    n53 = sum(len(s.refs_800_53) for s in catalog.subcategories)
    niso = sum(len(s.refs_iso_27001_2022) for s in catalog.subcategories)
    with_refs = sum(1 for s in catalog.subcategories if s.refs_800_53 or s.refs_iso_27001_2022)
    print(
        f"{len(catalog.functions)} functions, {len(catalog.categories)} categories, "
        f"{len(catalog.subcategories)} subcategories ({with_refs} with refs); "
        f"{n53} CSF->800-53 pairs, {niso} CSF->ISO 27001 pairs -> {args.out}"
    )


if __name__ == "__main__":
    main()
